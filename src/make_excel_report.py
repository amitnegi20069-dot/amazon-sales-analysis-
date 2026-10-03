"""
Builds reports/Amazon_Analysis_Summary.xlsx
- 'Products' sheet: the clean data
- Summary sheets use Excel formulas (COUNTIFS / AVERAGEIFS) that point to the Products sheet,
  so they update if the data is changed.

Usage: python src/make_excel_report.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
df = pd.read_csv(ROOT / "data" / "processed" / "amazon_products_clean.csv")
cols = ["product_id", "product_name", "brand", "main_category", "sub_category",
        "actual_price", "discounted_price", "discount_pct", "rating", "rating_count"]
df = df[cols]
n = len(df) + 1  # last row of data

FONT = "Arial"
head_font = Font(name=FONT, bold=True, color="FFFFFF")
head_fill = PatternFill("solid", fgColor="2F6690")
body_font = Font(name=FONT)

wb = Workbook()

# ---------- Products sheet ----------
ws = wb.active
ws.title = "Products"
ws.append(cols)
for row in df.itertuples(index=False):
    ws.append([None if pd.isna(v) else v for v in row])
for c in ws[1]:
    c.font, c.fill = head_font, head_fill
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.font = body_font
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"
for i, w in enumerate([14, 55, 14, 22, 30, 12, 14, 12, 8, 13], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# column letters in Products sheet
MAIN, SUB, DISC, PRICE, RATE, RC = "D", "E", "H", "G", "I", "J"
rng = lambda col: f"Products!${col}$2:${col}${n}"


def style_header(sheet, row, ncols):
    for c in range(1, ncols + 1):
        cell = sheet.cell(row=row, column=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


# ---------- Category summary ----------
cs = wb.create_sheet("Category Summary")
cs.append(["Main category", "Products", "Avg discount %", "Avg rating", "Total ratings"])
style_header(cs, 1, 5)
cats = df["main_category"].value_counts().index.tolist()
for i, cat in enumerate(cats, start=2):
    cs.cell(i, 1, cat)
    cs.cell(i, 2, f"=COUNTIFS({rng(MAIN)},A{i})")
    cs.cell(i, 3, f"=AVERAGEIFS({rng(DISC)},{rng(MAIN)},A{i})")
    cs.cell(i, 4, f"=AVERAGEIFS({rng(RATE)},{rng(MAIN)},A{i})")
    cs.cell(i, 5, f"=SUMIFS({rng(RC)},{rng(MAIN)},A{i})")
last = len(cats) + 1
tot = last + 1
cs.cell(tot, 1, "Total / overall")
cs.cell(tot, 2, f"=SUM(B2:B{last})")
cs.cell(tot, 3, f"=AVERAGE({rng(DISC)})")
cs.cell(tot, 4, f"=AVERAGE({rng(RATE)})")
cs.cell(tot, 5, f"=SUM(E2:E{last})")
for r in range(2, tot + 1):
    for c in range(1, 6):
        cs.cell(r, c).font = Font(name=FONT, bold=(r == tot))
    cs.cell(r, 2).number_format = "#,##0"
    cs.cell(r, 3).number_format = "0.0"
    cs.cell(r, 4).number_format = "0.00"
    cs.cell(r, 5).number_format = "#,##0"
cs.cell(tot + 2, 1, "Note: the last 5 categories have 2 or fewer products each, so their averages are not meaningful.").font = Font(name=FONT, italic=True)
cs.column_dimensions["A"].width = 26
for col in "BCDE":
    cs.column_dimensions[col].width = 16

# ---------- Sub-category summary ----------
ss = wb.create_sheet("Subcategory Summary")
ss.append(["Sub-category", "Products", "Avg discount %", "Avg rating", "Total ratings"])
style_header(ss, 1, 5)
sub_counts = df["sub_category"].value_counts()
subs = sub_counts[sub_counts >= 15].index.tolist()
for i, s in enumerate(subs, start=2):
    ss.cell(i, 1, s)
    ss.cell(i, 2, f"=COUNTIFS({rng(SUB)},A{i})")
    ss.cell(i, 3, f"=AVERAGEIFS({rng(DISC)},{rng(SUB)},A{i})")
    ss.cell(i, 4, f"=AVERAGEIFS({rng(RATE)},{rng(SUB)},A{i})")
    ss.cell(i, 5, f"=SUMIFS({rng(RC)},{rng(SUB)},A{i})")
    for c in range(1, 6):
        ss.cell(i, c).font = body_font
    ss.cell(i, 3).number_format = "0.0"
    ss.cell(i, 4).number_format = "0.00"
    ss.cell(i, 5).number_format = "#,##0"
ss.cell(len(subs) + 3, 1, "Only sub-categories with 15 or more products are listed.").font = Font(name=FONT, italic=True)
ss.column_dimensions["A"].width = 34
for col in "BCDE":
    ss.column_dimensions[col].width = 16

# ---------- Discount bands ----------
db = wb.create_sheet("Discount Bands")
db.append(["Band", "Min discount %", "Max discount %", "Products", "Avg rating"])
style_header(db, 1, 5)
bands = [("0-20%", 0, 20), ("21-40%", 21, 40), ("41-60%", 41, 60), ("61-80%", 61, 80), ("81%+", 81, 100)]
for i, (name, lo, hi) in enumerate(bands, start=2):
    db.cell(i, 1, name)
    db.cell(i, 2, lo)
    db.cell(i, 3, hi)
    db.cell(i, 4, f'=COUNTIFS({rng(DISC)},">="&B{i},{rng(DISC)},"<="&C{i})')
    db.cell(i, 5, f'=AVERAGEIFS({rng(RATE)},{rng(DISC)},">="&B{i},{rng(DISC)},"<="&C{i})')
    for c in range(1, 6):
        db.cell(i, c).font = body_font
    db.cell(i, 5).number_format = "0.00"
db.cell(8, 1, "Min and max columns are the band limits (inputs). Change them and the table updates.").font = Font(name=FONT, italic=True)
db.column_dimensions["A"].width = 12
for col in "BCDE":
    db.column_dimensions[col].width = 16

# ---------- Price bands ----------
pb = wb.create_sheet("Price Bands")
pb.append(["Band", "Min price (Rs)", "Max price (Rs)", "Products", "Avg discount %", "Avg rating"])
style_header(pb, 1, 6)
# a price belongs to a band if: min < price <= max (same rule as the pandas notebooks)
pbands = [("Under 200", 0, 200), ("200-500", 200, 500), ("500-1000", 500, 1000),
          ("1000-2000", 1000, 2000), ("2000-5000", 2000, 5000), ("5000+", 5000, 10000000)]
for i, (name, lo, hi) in enumerate(pbands, start=2):
    pb.cell(i, 1, name)
    pb.cell(i, 2, lo)
    pb.cell(i, 3, hi)
    cond = f'{rng(PRICE)},">"&B{i},{rng(PRICE)},"<="&C{i}'
    pb.cell(i, 4, f"=COUNTIFS({cond})")
    pb.cell(i, 5, f"=AVERAGEIFS({rng(DISC)},{cond})")
    pb.cell(i, 6, f"=AVERAGEIFS({rng(RATE)},{cond})")
    for c in range(1, 7):
        pb.cell(i, c).font = body_font
    pb.cell(i, 5).number_format = "0.0"
    pb.cell(i, 6).number_format = "0.00"
    pb.cell(i, 2).number_format = "#,##0"
    pb.cell(i, 3).number_format = "#,##0"
pb.cell(8, 1, "A price is in a band if min < price <= max.").font = Font(name=FONT, italic=True)
pb.column_dimensions["A"].width = 14
for col in "BCDEF":
    pb.column_dimensions[col].width = 16

wb.move_sheet("Products", offset=len(wb.sheetnames) - 1)
wb.active = 0
out = ROOT / "reports" / "Amazon_Analysis_Summary.xlsx"
wb.save(out)
print("Saved", out)
